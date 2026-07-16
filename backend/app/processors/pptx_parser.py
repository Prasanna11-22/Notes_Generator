"""
PPTX Parser.

Extracts text from Microsoft PowerPoint presentations using ``python-pptx``.
Text is gathered from every slide's text frames (both body and title
placeholders) and from the speaker-notes section.
"""

from __future__ import annotations

import io

from loguru import logger

from app.processors.base import BaseParser, ParseResult

_SUPPORTED_MIME_TYPES = frozenset(
    {
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }
)


class PPTXParser(BaseParser):
    """Extract text and metadata from ``.pptx`` files via python-pptx."""

    @property
    def supported_mime_types(self) -> frozenset[str]:
        return _SUPPORTED_MIME_TYPES

    def parse(self, file_bytes: bytes, filename: str) -> ParseResult:
        """Extract text from a PPTX byte stream."""
        logger.debug("PPTXParser: parsing '{}' ({} bytes)", filename, len(file_bytes))

        from pptx import Presentation  # type: ignore[import-untyped]
        from pptx.exc import PackageNotFoundError  # type: ignore[import-untyped]

        try:
            prs = Presentation(io.BytesIO(file_bytes))
        except (PackageNotFoundError, Exception) as exc:
            raise RuntimeError(
                f"'{filename}' is not a valid PPTX file: {exc}"
            ) from exc

        slide_parts: list[str] = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_texts: list[str] = []

            # ── Slide body text frames ────────────────────────────────────────
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        slide_texts.append(text)

            # ── Speaker notes ─────────────────────────────────────────────────
            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                notes_text = notes_frame.text.strip()
                if notes_text:
                    slide_texts.append(f"[Notes] {notes_text}")

            if slide_texts:
                slide_parts.append(
                    f"--- Slide {slide_num} ---\n" + "\n".join(slide_texts)
                )

        raw_text = "\n\n".join(slide_parts)
        page_count = len(prs.slides)

        # ── Core properties ───────────────────────────────────────────────────
        metadata: dict = {
            "page_count": page_count,
            "parser_backend": "python-pptx",
        }
        try:
            props = prs.core_properties
            if props.title:
                metadata["title"] = props.title
            if props.author:
                metadata["author"] = props.author
            if props.created:
                metadata["created"] = props.created.isoformat()
            if props.modified:
                metadata["modified"] = props.modified.isoformat()
        except Exception:
            pass

        return ParseResult(
            raw_text=raw_text,
            page_count=page_count,
            parser_name="PPTXParser(python-pptx)",
            metadata=metadata,
            is_scanned=False,
        )
