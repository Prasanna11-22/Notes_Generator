"""
DOCX Parser.

Extracts text from Microsoft Word documents using ``python-docx``.
Paragraph text, table cell contents, and document core properties are
all captured.
"""

from __future__ import annotations

import io

from loguru import logger

from app.processors.base import BaseParser, ParseResult

_SUPPORTED_MIME_TYPES = frozenset(
    {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
)


class DOCXParser(BaseParser):
    """Extract text and metadata from ``.docx`` files via python-docx."""

    @property
    def supported_mime_types(self) -> frozenset[str]:
        return _SUPPORTED_MIME_TYPES

    def parse(self, file_bytes: bytes, filename: str) -> ParseResult:
        """Extract text from a DOCX byte stream."""
        logger.debug("DOCXParser: parsing '{}' ({} bytes)", filename, len(file_bytes))

        from docx import Document  # type: ignore[import-untyped]

        try:
            doc = Document(io.BytesIO(file_bytes))
        except Exception as exc:
            raise RuntimeError(
                f"'{filename}' is not a valid DOCX file: {exc}"
            ) from exc

        parts: list[str] = []

        # ── Paragraph text ────────────────────────────────────────────────────
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        # ── Table cell text ───────────────────────────────────────────────────
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    parts.append(" | ".join(row_cells))

        raw_text = "\n\n".join(parts)

        # ── Core properties (title, author, dates) ────────────────────────────
        metadata: dict = {"parser_backend": "python-docx"}
        try:
            props = doc.core_properties
            if props.title:
                metadata["title"] = props.title
            if props.author:
                metadata["author"] = props.author
            if props.created:
                metadata["created"] = props.created.isoformat()
            if props.modified:
                metadata["modified"] = props.modified.isoformat()
        except Exception:
            pass  # core properties are optional

        # DOCX files don't have a meaningful "page count" without rendering;
        # we approximate it from paragraph count for informational purposes only.
        approx_page_count = max(1, len(doc.paragraphs) // 30)
        metadata["approx_page_count"] = approx_page_count

        return ParseResult(
            raw_text=raw_text,
            page_count=approx_page_count,
            parser_name="DOCXParser(python-docx)",
            metadata=metadata,
            is_scanned=False,
        )
